"""Text extraction with light structural context (section/page) so chunks can later carry
a section/clause label for citation (FR-005–FR-008)."""
import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

import docx
import pymupdf as fitz

_HEADING_PATTERN = re.compile(r"^(?:[A-Z][A-Za-z0-9 /&,-]{2,60}|[0-9]+(?:\.[0-9]+)*\s+.{3,60})$")


@dataclass
class Section:
    label: str
    page: int | None
    text: str


@dataclass
class ExtractedDocument:
    sections: list[Section] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(s.text for s in self.sections)


def _looks_like_heading(line: str) -> bool:
    stripped = line.strip()
    if not stripped or len(stripped) > 70 or stripped.endswith((".", ",", ";")):
        return False
    return bool(_HEADING_PATTERN.match(stripped))


def _group_lines_into_sections(lines: list[str], page: int | None) -> list[Section]:
    sections: list[Section] = []
    current_label = "General"
    current_lines: list[str] = []

    def flush():
        text = "\n".join(current_lines).strip()
        if text:
            sections.append(Section(label=current_label, page=page, text=text))

    for line in lines:
        if _looks_like_heading(line):
            flush()
            current_label = line.strip()
            current_lines = []
        else:
            current_lines.append(line)
    flush()
    return sections


def extract_pdf(file_path: str) -> ExtractedDocument:
    doc = ExtractedDocument()
    with fitz.open(file_path) as pdf:
        for page_index, page in enumerate(pdf, start=1):
            lines = page.get_text().splitlines()
            doc.sections.extend(_group_lines_into_sections(lines, page_index))
    return doc


def extract_docx(file_path: str) -> ExtractedDocument:
    document = docx.Document(file_path)
    doc = ExtractedDocument()
    current_label = "General"
    current_lines: list[str] = []

    def flush():
        text = "\n".join(current_lines).strip()
        if text:
            doc.sections.append(Section(label=current_label, page=None, text=text))

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name and paragraph.style.name.startswith("Heading"):
            flush()
            current_label = text
            current_lines = []
        else:
            current_lines.append(text)
    flush()
    return doc


def extract_txt(file_path: str) -> ExtractedDocument:
    with open(file_path, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()
    doc = ExtractedDocument()
    doc.sections = _group_lines_into_sections(lines, None)
    return doc


def extract_csv(file_path: str) -> ExtractedDocument:
    with open(file_path, "r", encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return ExtractedDocument()
    header, *data_rows = rows
    lines = [", ".join(f"{h}: {v}" for h, v in zip(header, row)) for row in data_rows]
    return ExtractedDocument(sections=[Section(label="General", page=None, text="\n".join(lines))])


def extract_text(file_path: str, filename: str) -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(file_path)
    if suffix == ".docx":
        return extract_docx(file_path)
    if suffix == ".csv":
        return extract_csv(file_path)
    return extract_txt(file_path)
